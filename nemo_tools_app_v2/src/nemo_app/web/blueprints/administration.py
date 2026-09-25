from __future__ import annotations

import uuid
from io import BytesIO

from flask import Blueprint, flash, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..common import checkbox, enqueue_upload_job, job_redirect, job_store

administration_blueprint = Blueprint("administration", __name__)


@administration_blueprint.get("/tools/user-import")
def import_form():
    return render_template("user_import_form.html")


@administration_blueprint.post("/tools/user-import")
def import_submit():
    upload = request.files.get("spreadsheet")
    token = request.form.get("api_token", "").strip()
    if not upload or not upload.filename or not token:
        flash("Choose a spreadsheet and enter your NEMO API token.", "error")
        return render_template("user_import_form.html"), 400
    suffix = ".csv" if upload.filename.lower().endswith(".csv") else ".xlsx"
    job_id = enqueue_upload_job(
        "user_import",
        title="User/account/project import",
        upload_specs=[(upload, {".csv", ".xlsx"})],
        payload={
            "input": f"input_1{suffix}",
            "dry_run": checkbox("dry_run"),
            "use_cache": not checkbox("bypass_cache"),
        },
        secrets={"api_token": token},
    )
    return job_redirect(job_id)


@administration_blueprint.get("/tools/qualification-import")
def qualification_import_form():
    return render_template("qualification_import_form.html")


@administration_blueprint.get("/tools/qualification-import/template.xlsx")
def qualification_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Qualifications"
    headers = ["Notes (optional)", "Tool ID", "Qualification Date", "User Email"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 34
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="NEMO-qualification-import-template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@administration_blueprint.post("/tools/qualification-import")
def qualification_import_submit():
    upload = request.files.get("spreadsheet")
    token = request.form.get("api_token", "").strip()
    if not upload or not upload.filename or not token:
        flash("Choose an Excel spreadsheet and enter your NEMO API token.", "error")
        return render_template("qualification_import_form.html"), 400
    job_id = enqueue_upload_job(
        "qualification_import",
        title="Qualification batch import",
        upload_specs=[(upload, {".xlsx"})],
        payload={
            "input": "input_1.xlsx",
            "dry_run": checkbox("dry_run"),
        },
        secrets={"api_token": token},
    )
    return job_redirect(job_id)


@administration_blueprint.get("/tools/replacement")
def replacement_form():
    return render_template("replacement_form.html")


@administration_blueprint.post("/tools/replacement")
def replacement_submit():
    token = request.form.get("api_token", "").strip()
    old_value = request.form.get("old_value", "").strip()
    target_value = request.form.get("target_value", "").strip()
    if not token or not old_value or not target_value:
        flash("Token, old value, and target value are required.", "error")
        return render_template("replacement_form.html"), 400
    job_id = str(uuid.uuid4())
    job_store().enqueue(
        "replacement",
        title="Account/project replacement",
        payload={
            "old_value": old_value,
            "target_value": target_value,
            "mode": request.form.get("mode", "new"),
            "dry_run": checkbox("dry_run"),
            "deactivate_old": checkbox("deactivate_old"),
        },
        secrets={"api_token": token},
        job_id=job_id,
    )
    return job_redirect(job_id)
