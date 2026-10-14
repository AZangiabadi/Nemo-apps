from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from nemo_app.billing.gold_audit import GoldAuditReport

HEADER_FILL = PatternFill("solid", fgColor="8A6700")
TITLE_FILL = PatternFill("solid", fgColor="FFF2CC")
MATCH_FILL = PatternFill("solid", fgColor="E2F0D9")
MISMATCH_FILL = PatternFill("solid", fgColor="F4CCCC")
MISSING_FILL = PatternFill("solid", fgColor="FCE5CD")
THIN_GRAY = Side(style="thin", color="D9D9D9")


def render_gold_audit_workbook(
    report: GoldAuditReport,
    output_path: Path,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gold Audit"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A11"
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = "10:10"

    sheet["A1"] = "Angstrom Gold Deposition Self-Reporting Audit"
    sheet["A1"].font = Font(size=16, bold=True, color="5B4500")
    sheet["A1"].fill = TITLE_FILL
    sheet["A2"] = (
        "Compares Au thickness entered in NEMO usage events with indexed instrument-log records. "
        "Review every highlighted exception before relying on the totals."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")

    sheet["A4"] = "Summary"
    sheet["A4"].font = Font(bold=True, color="FFFFFF")
    sheet["A4"].fill = HEADER_FILL
    summary_labels = [
        "User-reported Au total (nm)",
        "Instrument-log Au total (nm)",
        "Difference: reported - log (nm)",
        "Exception rows",
    ]
    for row, label in enumerate(summary_labels, 5):
        sheet.cell(row, 1, label)

    sheet["P1"] = "Audit rules"
    sheet["P1"].font = Font(bold=True, color="FFFFFF")
    sheet["P1"].fill = HEADER_FILL
    rules = [
        ("Log source", report.source_status),
        ("Time allowance (minutes)", report.match_grace_minutes),
        ("Absolute tolerance (nm)", report.absolute_tolerance_nm),
        ("Relative tolerance", report.relative_tolerance),
        ("Instrument mapping", "High Vacuum → UHV; Metals Deposition System → Classic"),
        ("Matching method", "Each log run is assigned to the closest same-instrument usage event"),
    ]
    for row, (label, value) in enumerate(rules, 2):
        sheet.cell(row, 16, label).font = Font(bold=True)
        sheet.cell(row, 17, value)
    sheet["Q5"].number_format = "0.0%"

    headers = [
        "Usage Start",
        "Usage End",
        "Instrument",
        "User",
        "Project",
        "Usage Event ID",
        "User Reported Au (nm)",
        "Log Au (nm)",
        "Difference (nm)",
        "Status",
        "Closest Log Date/Time",
        "Time Difference (min)",
        "Log Run ID(s)",
        "Log Source(s)",
        "Notes",
    ]
    header_row = 10
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)

    first_data_row = header_row + 1
    for row_index, record in enumerate(report.rows, first_data_row):
        values = [
            record.usage_start,
            record.usage_end,
            record.instrument,
            record.user,
            record.project,
            record.event_id,
            record.reported_nm,
            record.logged_nm,
            None,
            None,
            record.closest_log_time,
            record.time_difference_minutes,
            record.log_run_ids,
            record.log_sources,
            record.notes,
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=column in {13, 14, 15})
            cell.border = Border(bottom=THIN_GRAY)
        sheet.cell(
            row_index, 9, f'=IF(OR(G{row_index}="",H{row_index}=""),"",G{row_index}-H{row_index})'
        )
        sheet.cell(
            row_index,
            10,
            f'=IF(G{row_index}="","No user report",IF(H{row_index}="","No log match",'
            f'IF(ABS(I{row_index})<=MAX($Q$4,$Q$5*ABS(H{row_index})),"Match","Mismatch")))',
        )

    last_data_row = max(first_data_row, first_data_row + len(report.rows) - 1)
    total_row = last_data_row + 1
    sheet.cell(total_row, 1, "Cumulative totals")
    sheet.cell(total_row, 1).font = Font(bold=True)
    sheet.cell(total_row, 7, f"=SUM(G{first_data_row}:G{last_data_row})")
    sheet.cell(total_row, 8, f"=SUM(H{first_data_row}:H{last_data_row})")
    sheet.cell(total_row, 9, f"=G{total_row}-H{total_row}")
    for column in range(1, len(headers) + 1):
        sheet.cell(total_row, column).fill = TITLE_FILL
        sheet.cell(total_row, column).border = Border(top=Side(style="medium", color="8A6700"))

    sheet["B5"] = f"=G{total_row}"
    sheet["B6"] = f"=H{total_row}"
    sheet["B7"] = f"=I{total_row}"
    sheet["B8"] = (
        f'=COUNTIF(J{first_data_row}:J{last_data_row},"<>Match")'
        f"-COUNTBLANK(J{first_data_row}:J{last_data_row})"
    )

    for coordinate in ("B5", "B6", "B7", f"G{total_row}", f"H{total_row}", f"I{total_row}"):
        sheet[coordinate].number_format = "#,##0.00"
    sheet["B8"].number_format = "#,##0"
    for row in range(first_data_row, last_data_row + 1):
        for column in (1, 2, 11):
            sheet.cell(row, column).number_format = "yyyy-mm-dd hh:mm:ss"
        for column in (7, 8, 9, 12):
            sheet.cell(row, column).number_format = "#,##0.00"

    if report.rows:
        table = Table(displayName="GoldAuditTable", ref=f"A{header_row}:O{last_data_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        audit_range = f"A{first_data_row}:O{last_data_row}"
        sheet.conditional_formatting.add(
            audit_range,
            FormulaRule(formula=[f'$J{first_data_row}="Match"'], fill=MATCH_FILL),
        )
        sheet.conditional_formatting.add(
            audit_range,
            FormulaRule(formula=[f'$J{first_data_row}="Mismatch"'], fill=MISMATCH_FILL),
        )
        sheet.conditional_formatting.add(
            audit_range,
            FormulaRule(
                formula=[
                    f'OR($J{first_data_row}="No user report",$J{first_data_row}="No log match")'
                ],
                fill=MISSING_FILL,
            ),
        )

    widths = {
        "A": 20,
        "B": 20,
        "C": 34,
        "D": 22,
        "E": 20,
        "F": 15,
        "G": 22,
        "H": 16,
        "I": 17,
        "J": 17,
        "K": 22,
        "L": 20,
        "M": 28,
        "N": 55,
        "O": 36,
        "P": 26,
        "Q": 64,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[2].height = 42
    sheet.row_dimensions[header_row].height = 38
    # The Excel table already owns the AutoFilter for this range.  Writing a
    # second worksheet-level AutoFilter over the same cells produces duplicate
    # filter records that desktop Excel may remove during workbook recovery.
    # For an empty report there is no table (and no data to filter), so no
    # worksheet-level AutoFilter is needed either.
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)
    workbook.close()
    return output_path
