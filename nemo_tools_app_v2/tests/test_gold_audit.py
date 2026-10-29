from __future__ import annotations

import datetime as dt
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook

from nemo_app.angstrom.client import AngstromLogClient
from nemo_app.billing.gold_audit import (
    LoggedGoldRun,
    ReportedGoldUsage,
    billable_logged_thickness,
    build_gold_audit_report,
)
from nemo_app.invoices.gold_audit_renderer import render_gold_audit_workbook


class _Response:
    def __init__(self, payload, status_code=200):
        self.payload = payload
        self.status_code = status_code

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")

    def json(self):
        return self.payload


class GoldAuditTests(unittest.TestCase):
    def test_billing_uses_logs_even_without_user_gold_report(self):
        usage = self.usage(1, dt.datetime(2026, 7, 27, 9), dt.datetime(2026, 7, 27, 10), 0)
        report = build_gold_audit_report(
            [usage],
            [
                self.log(dt.datetime(2026, 7, 27, 9, 20), 40, "run-a"),
                self.log(dt.datetime(2026, 7, 27, 9, 40), 60, "run-b"),
            ],
        )
        self.assertEqual(billable_logged_thickness(report), {1: 100})

    def test_missing_and_ambiguous_logs_are_not_used_for_billing(self):
        usages = [
            self.usage(i, dt.datetime(2026, 7, 27, 9), dt.datetime(2026, 7, 27, 10), 50)
            for i in (1, 2)
        ]
        self.assertEqual(billable_logged_thickness(build_gold_audit_report(usages[:1], [])), {})
        self.assertEqual(
            billable_logged_thickness(
                build_gold_audit_report(
                    usages,
                    [self.log(dt.datetime(2026, 7, 27, 9, 20), 50, "run-a")],
                )
            ),
            {},
        )

    def usage(
        self,
        event_id: int,
        start: dt.datetime,
        end: dt.datetime,
        reported_nm: float,
    ) -> ReportedGoldUsage:
        return ReportedGoldUsage(
            event_id=event_id,
            system="Angstrom UHV",
            instrument="Angstrom High Vacuum",
            user="Ada Lovelace",
            project="Project 450",
            start=start,
            end=end,
            reported_nm=reported_nm,
        )

    def log(self, start: dt.datetime, thickness_nm: float, run_id: str) -> LoggedGoldRun:
        return LoggedGoldRun(
            system="Angstrom UHV",
            start=start,
            thickness_nm=thickness_nm,
            run_id=run_id,
            source="Pocket 1",
            source_file=f"uhv/{run_id}.csv",
            warning="",
        )

    def test_logs_are_grouped_into_the_closest_usage_session(self) -> None:
        usages = [
            self.usage(
                1,
                dt.datetime(2026, 7, 27, 9, 0),
                dt.datetime(2026, 7, 27, 10, 30),
                100.0,
            ),
            self.usage(
                2,
                dt.datetime(2026, 7, 27, 13, 0),
                dt.datetime(2026, 7, 27, 14, 0),
                50.0,
            ),
        ]
        logs = [
            self.log(dt.datetime(2026, 7, 27, 9, 20), 60.0, "run-a"),
            self.log(dt.datetime(2026, 7, 27, 10, 0), 40.0, "run-b"),
            self.log(dt.datetime(2026, 7, 27, 18, 0), 25.0, "run-unreported"),
        ]

        report = build_gold_audit_report(usages, logs)

        self.assertEqual(len(report.rows), 3)
        first = next(row for row in report.rows if row.event_id == 1)
        self.assertEqual(first.logged_nm, 100.0)
        self.assertEqual(first.initial_status, "Match")
        self.assertEqual(first.log_run_ids, "run-a; run-b")
        second = next(row for row in report.rows if row.event_id == 2)
        self.assertEqual(second.initial_status, "No log match")
        unreported = next(row for row in report.rows if row.event_id is None)
        self.assertEqual(unreported.logged_nm, 25.0)
        self.assertEqual(unreported.initial_status, "No user report")
        self.assertEqual(report.exception_count, 2)

    def test_tolerance_uses_larger_of_one_nm_or_five_percent(self) -> None:
        usage = self.usage(
            1,
            dt.datetime(2026, 7, 27, 9, 0),
            dt.datetime(2026, 7, 27, 10, 0),
            104.0,
        )
        matched = build_gold_audit_report(
            [usage], [self.log(dt.datetime(2026, 7, 27, 9, 10), 100.0, "run-a")]
        )
        mismatched = build_gold_audit_report(
            [usage], [self.log(dt.datetime(2026, 7, 27, 9, 10), 98.0, "run-b")]
        )
        self.assertEqual(matched.rows[0].initial_status, "Match")
        self.assertEqual(mismatched.rows[0].initial_status, "Mismatch")

    def test_workbook_contains_formula_totals_and_exception_highlighting(self) -> None:
        report = build_gold_audit_report(
            [
                self.usage(
                    1,
                    dt.datetime(2026, 7, 27, 9, 0),
                    dt.datetime(2026, 7, 27, 10, 0),
                    80.0,
                )
            ],
            [self.log(dt.datetime(2026, 7, 27, 9, 10), 100.0, "run-a")],
        )
        with tempfile.TemporaryDirectory() as folder:
            path = render_gold_audit_workbook(report, Path(folder) / "audit.xlsx")
            workbook = load_workbook(path, data_only=False)
            sheet = workbook["Gold Audit"]
            self.assertEqual(sheet["G11"].value, 80.0)
            self.assertEqual(sheet["H11"].value, 100.0)
            self.assertEqual(sheet["I11"].value, '=IF(OR(G11="",H11=""),"",G11-H11)')
            self.assertIn('"Mismatch"', sheet["J11"].value)
            self.assertEqual(sheet["G12"].value, "=SUM(G11:G11)")
            self.assertEqual(sheet["H12"].value, "=SUM(H11:H11)")
            self.assertEqual(len(sheet.conditional_formatting), 1)
            self.assertIn("GoldAuditTable", sheet.tables)
            workbook.close()

            with zipfile.ZipFile(path) as archive:
                worksheet_xml = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                table_xml = ElementTree.fromstring(archive.read("xl/tables/table1.xml"))

            spreadsheet_ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            self.assertIsNone(worksheet_xml.find("main:autoFilter", spreadsheet_ns))
            self.assertEqual(
                table_xml.find("main:autoFilter", spreadsheet_ns).attrib["ref"],
                "A10:O11",
            )

    def test_log_client_uses_existing_async_scan_api(self) -> None:
        client = AngstromLogClient(
            "https://angstrom.example",
            poll_interval=0,
            max_wait=1,
        )
        requests_seen = []
        states = iter(
            [
                {"state": "running", "percent": 60},
                {"state": "complete", "result": {"runs": [{"run_id": "one"}]}},
            ]
        )

        def post(url, *, json, timeout):
            requests_seen.append(("POST", url, json, timeout))
            return _Response({"job_id": "abc"})

        def get(url, *, timeout):
            requests_seen.append(("GET", url, None, timeout))
            return _Response(next(states))

        client.session.post = post  # type: ignore[method-assign]
        client.session.get = get  # type: ignore[method-assign]

        runs = client.fetch_gold_runs(dt.date(2026, 7, 1), dt.date(2026, 7, 31))

        self.assertEqual(runs, [{"run_id": "one"}])
        self.assertEqual(requests_seen[0][1], "https://angstrom.example/api/scans")
        self.assertEqual(
            requests_seen[0][2],
            {"start": "2026-07-01", "end": "2026-07-31"},
        )
        self.assertEqual(requests_seen[-1][1], "https://angstrom.example/api/scans/abc")


if __name__ == "__main__":
    unittest.main()
